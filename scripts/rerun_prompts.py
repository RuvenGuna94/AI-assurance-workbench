from __future__ import annotations

import argparse
import csv
import os
import time
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import httpx
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

EXTRA_COLUMNS = [
    "Rerun_ID",
    "Rerun_Number",
    "Rerun_Timestamp_UTC",
    "HTTP_Status",
    "Request_Succeeded",
    "Latency_ms",
    "Response_ID",
    "Response_Model",
    "Rerun_Output",
    "Request_Error",
    "Reproduced",
    "Reviewer_Classification",
    "Rerun_Reviewer_Notes",
]


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Rerun selected prompts against the local support gateway."
    )
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--runs-per-prompt", type=int, default=3)
    parser.add_argument(
        "--base-url",
        default=os.getenv("GATEWAY_BASE_URL", "http://localhost:8000"),
    )
    parser.add_argument(
        "--timeout",
        type=float,
        default=float(os.getenv("REQUEST_TIMEOUT_SECONDS", "120")),
    )
    return parser.parse_args()


def excel_safe_text(value: Any) -> Any:
    """Prevent untrusted prompt text from being interpreted as an Excel formula."""
    if value is None:
        return ""

    if isinstance(value, (bool, int, float)):
        return value

    text = str(value)

    if text.startswith(("=", "+", "-", "@")):
        return "'" + text

    return text


def extract_output(response_data: dict[str, Any]) -> str:
    choices = response_data.get("choices") or []

    if not choices:
        return ""

    first_choice = choices[0] or {}
    message = first_choice.get("message") or {}
    return str(message.get("content") or "")


def read_source_rows(
    input_path: Path,
) -> tuple[list[str], list[dict[str, str]]]:
    with input_path.open("r", encoding="utf-8-sig", newline="") as source_file:
        reader = csv.DictReader(source_file)
        source_columns = reader.fieldnames or []
        rows = list(reader)

    if not rows:
        raise ValueError("The input CSV contains no data rows.")

    if "Prompt" not in source_columns:
        raise ValueError("The input CSV must contain a Prompt column.")

    blank_rows = [
        index
        for index, row in enumerate(rows, start=2)
        if not str(row.get("Prompt") or "").strip()
    ]

    if blank_rows:
        raise ValueError(f"Blank prompts were found on CSV rows: {blank_rows}")

    return source_columns, rows


def run_prompts(
    source_rows: list[dict[str, str]],
    runs_per_prompt: int,
    base_url: str,
    api_key: str,
    timeout: float,
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    endpoint = f"{base_url.rstrip('/')}/v1/chat/completions"

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }

    with httpx.Client(timeout=timeout) as client:
        for source_row in source_rows:
            prompt = str(source_row["Prompt"]).strip()
            attempt_uuid = str(source_row.get("Attempt_UUID") or "unknown")

            for rerun_number in range(1, runs_per_prompt + 1):
                timestamp = datetime.now(UTC).isoformat()
                started = time.perf_counter()

                status_code: int | str = ""
                request_succeeded = False
                response_id = ""
                response_model = ""
                rerun_output = ""
                request_error = ""

                payload = {
                    "model": "northstar-support",
                    "messages": [
                        {
                            "role": "user",
                            "content": prompt,
                        }
                    ],
                    "temperature": 0.0,
                    "max_tokens": 300,
                }

                try:
                    response = client.post(
                        endpoint,
                        headers=headers,
                        json=payload,
                    )
                    status_code = response.status_code

                    try:
                        response_data = response.json()
                    except ValueError:
                        response_data = {}

                    if response.is_success:
                        rerun_output = extract_output(response_data)
                        response_id = str(response_data.get("id") or "")
                        response_model = str(response_data.get("model") or "")

                        if rerun_output:
                            request_succeeded = True
                        else:
                            request_error = (
                                "Successful HTTP response did not contain "
                                "choices[0].message.content."
                            )
                    else:
                        request_error = response.text[:2000]

                except httpx.HTTPError as error:
                    request_error = (f"{type(error).__name__}: {error}")[:2000]

                latency_ms = round(
                    (time.perf_counter() - started) * 1000,
                    1,
                )

                result: dict[str, Any] = dict(source_row)
                result.update(
                    {
                        "Rerun_ID": (f"{attempt_uuid}-rerun-{rerun_number}"),
                        "Rerun_Number": rerun_number,
                        "Rerun_Timestamp_UTC": timestamp,
                        "HTTP_Status": status_code,
                        "Request_Succeeded": request_succeeded,
                        "Latency_ms": latency_ms,
                        "Response_ID": response_id,
                        "Response_Model": response_model,
                        "Rerun_Output": rerun_output,
                        "Request_Error": request_error,
                        "Reproduced": "",
                        "Reviewer_Classification": "",
                        "Rerun_Reviewer_Notes": "",
                    }
                )
                results.append(result)

                print(
                    f"{result['Rerun_ID']}: "
                    f"HTTP {status_code or 'ERROR'}, "
                    f"success={request_succeeded}, "
                    f"{latency_ms} ms"
                )

    return results


def create_workbook(
    output_path: Path,
    source_name: str,
    source_row_count: int,
    runs_per_prompt: int,
    base_url: str,
    source_columns: list[str],
    results: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    results_sheet = workbook.active
    results_sheet.title = "Results"
    results_sheet.sheet_view.showGridLines = False

    output_columns = source_columns + EXTRA_COLUMNS
    results_sheet.append(output_columns)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in results_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for result in results:
        results_sheet.append(
            [excel_safe_text(result.get(column, "")) for column in output_columns]
        )

    results_sheet.freeze_panes = "A2"
    results_sheet.row_dimensions[1].height = 35

    long_text_columns = {
        "Prompt": 65,
        "Output": 65,
        "Rerun_Output": 65,
        "Request_Error": 45,
        "Reviewer_Notes": 40,
        "Rerun_Reviewer_Notes": 40,
        "Goal": 35,
        "Trigger": 35,
    }

    identifier_columns = {
        "Run_ID": 38,
        "Attempt_UUID": 38,
        "Rerun_ID": 50,
        "Probe": 36,
        "Detector": 36,
        "Response_ID": 38,
    }

    for index, column_name in enumerate(output_columns, start=1):
        column_letter = get_column_letter(index)

        if column_name in long_text_columns:
            width = long_text_columns[column_name]
        elif column_name in identifier_columns:
            width = identifier_columns[column_name]
        else:
            width = max(14, min(len(column_name) + 3, 25))

        results_sheet.column_dimensions[column_letter].width = width

        for cell in results_sheet[column_letter][1:]:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    final_row = results_sheet.max_row
    final_column = get_column_letter(results_sheet.max_column)

    table = Table(
        displayName="PromptRerunResults",
        ref=f"A1:{final_column}{final_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    results_sheet.add_table(table)

    reproduced_column = output_columns.index("Reproduced") + 1
    classification_column = output_columns.index("Reviewer_Classification") + 1
    success_column = output_columns.index("Request_Succeeded") + 1

    reproduced_validation = DataValidation(
        type="list",
        formula1='"Yes,No,Unclear"',
        allow_blank=True,
    )
    classification_validation = DataValidation(
        type="list",
        formula1=(
            '"Confirmed finding,Not reproduced,Inconclusive,Needs investigation"'
        ),
        allow_blank=True,
    )

    results_sheet.add_data_validation(reproduced_validation)
    results_sheet.add_data_validation(classification_validation)

    reproduced_validation.add(
        f"{get_column_letter(reproduced_column)}2:"
        f"{get_column_letter(reproduced_column)}{final_row}"
    )
    classification_validation.add(
        f"{get_column_letter(classification_column)}2:"
        f"{get_column_letter(classification_column)}{final_row}"
    )

    failed_fill = PatternFill("solid", fgColor="F4CCCC")
    success_letter = get_column_letter(success_column)

    results_sheet.conditional_formatting.add(
        f"{success_letter}2:{success_letter}{final_row}",
        FormulaRule(
            formula=[f"${success_letter}2=FALSE"],
            fill=failed_fill,
        ),
    )

    summary_sheet = workbook.create_sheet("Run Summary")
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.append(["Metric", "Value"])

    successful_requests = sum(1 for result in results if result["Request_Succeeded"])

    summary_rows = [
        ("Source file", source_name),
        ("Selected prompts", source_row_count),
        ("Runs per prompt", runs_per_prompt),
        ("Total requests", len(results)),
        ("Successful requests", successful_requests),
        (
            "Unsuccessful requests",
            len(results) - successful_requests,
        ),
        ("Gateway base URL", base_url),
        (
            "Workbook generated UTC",
            datetime.now(UTC).isoformat(),
        ),
    ]

    for row in summary_rows:
        summary_sheet.append(row)

    for cell in summary_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    summary_sheet.column_dimensions["A"].width = 28
    summary_sheet.column_dimensions["B"].width = 60
    summary_sheet.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    args = parse_arguments()

    if args.runs_per_prompt < 1:
        raise ValueError("--runs-per-prompt must be at least 1.")

    api_key = os.getenv("APP_API_KEY")

    if not api_key:
        raise RuntimeError("APP_API_KEY is not set in this PowerShell session.")

    source_columns, source_rows = read_source_rows(args.input)

    results = run_prompts(
        source_rows=source_rows,
        runs_per_prompt=args.runs_per_prompt,
        base_url=args.base_url,
        api_key=api_key,
        timeout=args.timeout,
    )

    create_workbook(
        output_path=args.output,
        source_name=args.input.name,
        source_row_count=len(source_rows),
        runs_per_prompt=args.runs_per_prompt,
        base_url=args.base_url,
        source_columns=source_columns,
        results=results,
    )

    print()
    print(f"Created workbook: {args.output}")
    print(f"Result rows: {len(results)}")


if __name__ == "__main__":
    main()
