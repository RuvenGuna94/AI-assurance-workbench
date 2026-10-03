"""Export Garak JSONL results to a review-friendly Excel workbook.

This exporter uses openpyxl and does not require Microsoft Excel to be
installed or licensed.
"""

from __future__ import annotations

import argparse
import json
import re
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ATTEMPT_HEADERS = [
    "Risk_ID",
    "Source_File",
    "Run_ID",
    "Probe",
    "Attempt_UUID",
    "Sequence",
    "Detector",
    "Detector_Score",
    "Review_Flag",
    "Prompt",
    "Output",
    "Trigger",
    "Goal",
    "Classification",
    "Finding_ID",
    "Reviewer_Notes",
]

SUMMARY_HEADERS = [
    "Risk_ID",
    "Source_File",
    "Run_ID",
    "Generations",
    "Probe",
    "Detector",
    "Passed",
    "Total",
    "Review_Signals",
    "Pass_Rate",
]

CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
FORMULA_PREFIXES = ("=", "+", "-", "@")
MAX_EXCEL_CELL_LENGTH = 32_767


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create an Excel review workbook from Garak JSONL output."
    )
    parser.add_argument(
        "--input",
        dest="inputs",
        action="append",
        required=True,
        type=Path,
        help="Garak JSONL input. Repeat this option to combine multiple runs.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("evidence/local/review/garak-review.xlsx"),
        help="Destination .xlsx path.",
    )
    parser.add_argument(
        "--review-threshold",
        type=float,
        default=0.5,
        help="Detector score at or above which a row is flagged for review.",
    )
    arguments = parser.parse_args()
    if not 0.0 <= arguments.review_threshold <= 1.0:
        parser.error("--review-threshold must be between 0 and 1")
    return arguments


def risk_id_for_probe(probe: str) -> str:
    if probe.startswith("promptinject."):
        return "R1"
    if probe.startswith("encoding."):
        return "R2"
    if "dan" in probe.lower() or probe.startswith("jailbreak."):
        return "R3"
    if "prompt" in probe.lower() and "leak" in probe.lower():
        return "R4"
    return ""


def to_review_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (list, tuple)):
        return "\n".join(to_review_text(item) for item in value)
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def safe_excel_text(value: Any) -> Any:
    """Make untrusted Garak text safe and valid inside an Excel cell."""
    if not isinstance(value, str):
        return value
    value = CONTROL_CHARACTERS.sub("", value)
    if value.startswith(FORMULA_PREFIXES):
        value = "'" + value
    if len(value) > MAX_EXCEL_CELL_LENGTH:
        suffix = "\n[truncated for Excel cell limit]"
        value = value[: MAX_EXCEL_CELL_LENGTH - len(suffix)] + suffix
    return value


def numeric_scores(value: Any) -> Iterable[float]:
    if isinstance(value, bool) or value is None:
        return
    if isinstance(value, (int, float)):
        yield float(value)
        return
    if isinstance(value, dict):
        for nested_value in value.values():
            yield from numeric_scores(nested_value)
        return
    if isinstance(value, (list, tuple)):
        for nested_value in value:
            yield from numeric_scores(nested_value)


def read_jsonl(path: Path) -> Iterable[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig") as source_file:
        for line_number, line in enumerate(source_file, start=1):
            if not line.strip():
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError as error:
                raise ValueError(
                    f"Invalid JSON in {path} at line {line_number}: {error}"
                ) from error
            if isinstance(record, dict):
                yield record


def parse_sources(
    paths: list[Path], review_threshold: float
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    attempt_rows: list[dict[str, Any]] = []
    summary_rows: list[dict[str, Any]] = []

    for path in paths:
        run_id = ""
        generations = 0
        for record in read_jsonl(path):
            entry_type = str(record.get("entry_type", ""))
            if entry_type == "start_run setup":
                generations = int(record.get("run.generations", 0) or 0)
                continue
            if entry_type == "init":
                run_id = str(record.get("run", ""))
                continue
            if entry_type == "attempt" and int(record.get("status", 0)) == 2:
                probe = str(record.get("probe_classname", ""))
                notes = record.get("notes")
                trigger = notes.get("triggers") if isinstance(notes, dict) else ""
                detector_results = record.get("detector_results", {})
                if not isinstance(detector_results, dict):
                    detector_results = {}
                for detector, raw_scores in detector_results.items():
                    scores = list(numeric_scores(raw_scores))
                    detector_score = max(scores, default=0.0)
                    attempt_rows.append(
                        {
                            "Risk_ID": risk_id_for_probe(probe),
                            "Source_File": path.name,
                            "Run_ID": run_id,
                            "Probe": probe,
                            "Attempt_UUID": str(record.get("uuid", "")),
                            "Sequence": int(record.get("seq", 0) or 0),
                            "Detector": str(detector),
                            "Detector_Score": detector_score,
                            "Review_Flag": (
                                "Review" if detector_score >= review_threshold else ""
                            ),
                            "Prompt": to_review_text(record.get("prompt")),
                            "Output": to_review_text(record.get("outputs")),
                            "Trigger": to_review_text(trigger),
                            "Goal": to_review_text(record.get("goal")),
                            "Classification": "",
                            "Finding_ID": "",
                            "Reviewer_Notes": "",
                        }
                    )
                continue
            if entry_type == "eval":
                passed = int(record.get("passed", 0) or 0)
                total = int(record.get("total", 0) or 0)
                probe = str(record.get("probe", ""))
                summary_rows.append(
                    {
                        "Risk_ID": risk_id_for_probe(probe),
                        "Source_File": path.name,
                        "Run_ID": run_id,
                        "Generations": generations,
                        "Probe": probe,
                        "Detector": str(record.get("detector", "")),
                        "Passed": passed,
                        "Total": total,
                        "Review_Signals": total - passed,
                        "Pass_Rate": passed / total if total else 0.0,
                    }
                )

    attempt_rows.sort(
        key=lambda row: (row["Risk_ID"], row["Probe"], row["Detector"], row["Sequence"])
    )
    summary_rows.sort(key=lambda row: (row["Risk_ID"], row["Probe"], row["Detector"]))
    return attempt_rows, summary_rows


def write_worksheet(
    worksheet: Any,
    rows: list[dict[str, Any]],
    headers: list[str],
    *,
    wrap_columns: set[int] | None = None,
    input_columns: set[int] | None = None,
    percent_columns: set[int] | None = None,
) -> None:
    wrap_columns = wrap_columns or set()
    input_columns = input_columns or set()
    percent_columns = percent_columns or set()
    header_fill = PatternFill("solid", fgColor="4F81BD")
    input_fill = PatternFill("solid", fgColor="FFF2CC")

    worksheet.append(headers)
    for row in rows:
        worksheet.append([safe_excel_text(row.get(header, "")) for header in headers])

    for cell in worksheet[1]:
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.row_dimensions[1].height = 30
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10)
            cell.alignment = Alignment(
                vertical="top", wrap_text=cell.column in wrap_columns
            )
            if cell.column in input_columns:
                cell.fill = input_fill
            if cell.column in percent_columns:
                cell.number_format = "0.00%"

    for column_number, header in enumerate(headers, start=1):
        if column_number in wrap_columns:
            width = 55
        else:
            longest_value = max(
                [len(header)] + [len(str(row.get(header, ""))) for row in rows[:500]]
            )
            width = min(max(longest_value + 2, 12), 35)
        worksheet.column_dimensions[get_column_letter(column_number)].width = width


def create_workbook(
    attempt_rows: list[dict[str, Any]],
    summary_rows: list[dict[str, Any]],
    output_path: Path,
) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Probe Summary"
    signals_sheet = workbook.create_sheet("Signals")
    attempts_sheet = workbook.create_sheet("All Attempts")
    signal_rows = [row for row in attempt_rows if row["Review_Flag"] == "Review"]

    write_worksheet(summary_sheet, summary_rows, SUMMARY_HEADERS, percent_columns={10})
    write_worksheet(
        signals_sheet,
        signal_rows,
        ATTEMPT_HEADERS,
        wrap_columns={10, 11, 12, 13, 16},
        input_columns={14, 15, 16},
    )
    write_worksheet(
        attempts_sheet,
        attempt_rows,
        ATTEMPT_HEADERS,
        wrap_columns={10, 11, 12, 13, 16},
        input_columns={14, 15, 16},
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_name(output_path.stem + ".tmp.xlsx")
    workbook.save(temporary_path)
    temporary_path.replace(output_path)

    print(f"Workbook created: {output_path.resolve()}")
    print(f"Summary rows: {len(summary_rows)}")
    print(f"Rows requiring review: {len(signal_rows)}")
    print(f"All detector rows: {len(attempt_rows)}")


def main() -> None:
    arguments = parse_arguments()
    input_paths = [path.resolve() for path in arguments.inputs]
    missing_paths = [path for path in input_paths if not path.is_file()]
    if missing_paths:
        missing = "\n".join(f"- {path}" for path in missing_paths)
        raise FileNotFoundError(f"Required input files were not found:\n{missing}")
    if arguments.output.suffix.lower() != ".xlsx":
        raise ValueError("The output filename must use the .xlsx extension")

    attempt_rows, summary_rows = parse_sources(input_paths, arguments.review_threshold)
    create_workbook(attempt_rows, summary_rows, arguments.output.resolve())


if __name__ == "__main__":
    main()
